"""Developer regressions on disposable fixtures. Not plugin acceptance evaluations."""
import argparse
import copy
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / 'scripts'))
import hwpdoc as h
from hwpdoc_config import Context, load_context, owned_workspace
import hwpdoc_templates as t
import hwpdoc_workspace as w
import build_distribution as packaging
import hwpdoc_setup as setup


class DistributionTests(unittest.TestCase):
    def setUp(self):
        scratch = REPO / '.claude/scratchpad/dev-tests'; scratch.mkdir(parents=True, exist_ok=True)
        temp = tempfile.TemporaryDirectory(dir=scratch); self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.settings = {'kind': 'hwpdoc-workspace', 'version': 1, 'school': {'document_prefix': 'TEST-'}, 'references': {}}
        h.write_json(self.root / '.hwpdoc/workspace.json', self.settings)
        self.ctx = Context(REPO, self.root, self.root / 'pc', self.settings, Path(sys.executable))
        for key, value in [('ROOT', self.root), ('JOBS', self.root / '.hwpdoc/jobs'), ('CONTEXT', self.ctx), ('ONEDRIVE', None)]:
            p = patch.object(h, key, value); p.start(); self.addCleanup(p.stop)
        self.source = self.root / 'source.hwpx'
        with ZipFile(self.source, 'w') as z:
            z.writestr('mimetype', 'application/hwp+zip')
            z.writestr('Contents/section0.xml', '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"><hp:p><hp:run><hp:t>TEST TITLE</hp:t></hp:run></hp:p></hs:sec>')

    def cli(self, *args):
        env = dict(os.environ, HWPDOC_WORKSPACE=str(self.root), HWPDOC_PC_DATA=str(self.root / 'pc'), PYTHONDONTWRITEBYTECODE='1')
        return subprocess.run([sys.executable, '-B', '-X', 'utf8', str(REPO / 'scripts/hwpdoc.py'), '--workspace', str(self.root), *args],
                              env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)

    def registration_args(self, **kwargs):
        return argparse.Namespace(id='test-form', version='1', source=None, mapping=None, review=None, purpose=['notice', 'reply'], provenance='native', **kwargs)

    def test_init_preserves_existing_instructions(self):
        target = self.root / 'another school'
        target.mkdir(); (target / 'AGENTS.md').write_text('KEEP', encoding='utf-8')
        args = argparse.Namespace(path=str(target), app='codex', skill_name='observed:hwpx', school=None, prefix=None, year=None, copy_to=None)
        with patch.object(h, 'configure'):
            result = w.init_workspace(args)
        self.assertEqual((target / 'AGENTS.md').read_text(), 'KEEP')
        self.assertEqual(result['mode'], 'existing_instructions_preserved')
        self.assertIn('references/workflow.md', Path(result['instructions']).read_text(encoding='utf-8'))
        self.assertEqual(result['actual_skill_loaded'], 'unverified')
        with self.assertRaises(ValueError): w.init_workspace(args)

    def test_init_never_creates_parent_agents(self):
        args = argparse.Namespace(path=str(REPO.parent))
        with self.assertRaises(ValueError): w.init_workspace(args)

    @unittest.skipUnless(os.name == 'nt', 'Windows hook launch')
    def test_hook_command_preserves_path_stdin_and_deny_across_shells(self):
        folder = self.root / "설치 공간 & '$`"
        folder.mkdir()
        script = folder / 'hook.ps1'
        script.write_text("param([string]$Event)\n[Console]::Out.Write([Console]::In.ReadToEnd()); exit 2\n", encoding='utf-8')
        command = setup.hook_command(script, 'pre', 'claude')
        shells = [['powershell.exe', '-NoProfile', '-NonInteractive', '-Command', setup.hook_command(script, 'pre', 'codex')]]
        bash = Path('C:/Program Files/Git/bin/bash.exe')
        if bash.exists(): shells.append([str(bash), '--noprofile', '--norc', '-c', command])
        payload = '{"tool_name":"Write","path":"synthetic"}'
        for argv in shells:
            with self.subTest(shell=argv[0]):
                result = subprocess.run(argv, input=payload, capture_output=True, text=True, encoding='utf-8', timeout=30)
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertEqual(result.stdout, payload)

    @unittest.skipUnless(os.name == 'nt', 'Windows path support boundary')
    def test_init_long_path_rejected_before_any_write(self):
        target = self.root / ('long-' + 'a' * 70) / ('b' * 70) / ('c' * 70)
        args = argparse.Namespace(path=str(target), app='codex')
        with patch.object(Path, 'mkdir') as mkdir:
            with self.assertRaisesRegex(ValueError, 'Windows 지원 범위.*더 짧은'):
                w.init_workspace(args)
            mkdir.assert_not_called()
        self.assertFalse(target.exists())

    @unittest.skipUnless(os.name == 'nt', 'Windows path support boundary')
    def test_init_checks_generated_children_not_only_root(self):
        target = self.root / ('a' * (235 - len(str(self.root)) - 1))
        self.assertLess(len(str(target)), 248)
        with self.assertRaisesRegex(ValueError, 'Windows 지원 범위'):
            w.check_init_paths(target)

    def test_context_move_and_pc_configuration(self):
        h.write_json(self.root / 'pc/runtime.json', {'version': 1, 'python': sys.executable})
        ctx = load_context(self.root, pc_data=self.root / 'pc')
        encoded = ctx.encode(self.source)
        second = self.root / 'moved'; second.mkdir()
        relocated = Context(REPO, second, ctx.pc_data, {}, ctx.python)
        self.assertEqual(relocated.decode(encoded), second / 'source.hwpx')
        self.assertEqual(ctx.python, Path(sys.executable).resolve())
        for path in ('../escape', str(REPO), 'workspace:../bad'):
            with self.subTest(path=path), self.assertRaises(ValueError):
                ctx.decode(path) if path.startswith('workspace:') else ctx.local(path)

    def test_job_junction_escape(self):
        outside = self.root / 'outside'; outside.mkdir()
        with patch.object(h, 'JOBS', self.root.parent):
            with self.assertRaises(ValueError): h.job_dir('escape')

    def test_malformed_settings_are_reported(self):
        for value in ([], {'kind': 'hwpdoc-workspace', 'version': 1, 'references': []}):
            h.write_json(self.root / '.hwpdoc/workspace.json', value)
            with self.assertRaises(ValueError): load_context(self.root, pc_data=self.root / 'pc')

    def test_init_conflicting_proposal_has_no_partial_marker(self):
        target = self.root / 'init-conflict'; (target / '.hwpdoc').mkdir(parents=True)
        (target / 'AGENTS.md').write_text('KEEP')
        (target / '.hwpdoc/instructions-to-merge.md').write_text('KEEP')
        args = argparse.Namespace(path=str(target), app='codex')
        with self.assertRaises(ValueError): w.init_workspace(args)
        self.assertFalse((target / '.hwpdoc/workspace.json').exists())

    def test_invalid_package_is_rejected_before_import(self):
        self.source.write_bytes(b'not a zip')
        args = self.registration_args(); args.source = str(self.source)
        with self.assertRaises(ValueError): t.add_template(args)
        self.assertFalse((t.folder_for(args.id, args.version) / 'original.hwpx').exists())

    def test_fresh_workspace_stop_checks_first_document(self):
        output = self.root / 'output'; output.mkdir()
        (output / 'broken.hwpx').write_bytes(b'invalid package')
        env = dict(os.environ, HWPDOC_WORKSPACE=str(self.root), HWPDOC_PC_DATA=str(self.root / 'pc'), PYTHONDONTWRITEBYTECODE='1')
        result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(REPO / 'scripts/validate_pipeline.py')],
                                env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        state = h.read_json(self.root / 'logs/.validate_state.json')
        self.assertFalse(state['broken.hwpx']['ok'])
        self.assertNotIn('baseline', state['broken.hwpx'])

    def test_timetable_mapping_uses_selected_school_cells(self):
        import openpyxl
        src = self.root / 'school.xlsx'
        workbook = openpyxl.Workbook(); sheet = workbook.active
        sheet.cell(1, 1, 'SYNTHETIC SCHOOL'); sheet.cell(4, 3, 'SCI'); sheet.cell(5, 3, 'MATH')
        workbook.save(src); workbook.close()
        settings = copy.deepcopy(self.settings)
        settings['references'] = {'timetable_source': 'school.xlsx', 'timetable': 'derived/table.json'}
        settings['timetable_mapping'] = dict(format='grade-columns-v1', grades={'class-z': 'Z'}, day_columns={'월': [3]},
                                             period_rows=[4, 5], subject_full={'SCI': 'Science', 'MATH': 'Math'}, non_subject=[])
        h.write_json(self.root / '.hwpdoc/workspace.json', settings)
        before = h.sha(src)
        env = dict(os.environ, HWPDOC_WORKSPACE=str(self.root), HWPDOC_PC_DATA=str(self.root / 'pc'), PYTHONDONTWRITEBYTECODE='1')
        result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(REPO / 'scripts/refresh_reference.py'), '--only', 'timetable'],
                                env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(h.read_json(self.root / 'derived/table.json')['시간표'], {'월': {'class-z': ['SCI', 'MATH']}})
        self.assertEqual(h.sha(src), before)

    def test_registration_receipt_detects_changed_purpose_and_source(self):
        folder = t.folder_for('receipt', '1'); folder.mkdir(parents=True)
        for name in t.FILES: (folder / name).write_text('SYNTHETIC FIXTURE')
        hashes = t.file_hashes(folder)
        review = dict(hashes=hashes, purposes=['notice'], provenance='native', confirmed=True,
                      by='SYNTHETIC_FIXTURE', at='2026-09-13', record='Test data, not real approval')
        h.write_json(folder / 'review.json', review)
        record = dict(version=1, status='registered', hashes=hashes, purposes=['notice'], provenance='native',
                      source={'sha256': hashes['original.hwpx']}, review_sha256=h.sha(folder / 'review.json'),
                      stages={s: {'status': 'pass'} for s in ('structure', 'namespaces', 'finalize', 'layout', 'profiles', 'hancom')})
        h.write_json(folder / 'registration.json', record)
        self.assertEqual(t.inspect_registration(folder)[1], [])
        record['purposes'].append('unreviewed'); record['source']['sha256'] = 'changed'
        h.write_json(folder / 'registration.json', record)
        self.assertEqual(len(t.inspect_registration(folder)[1]), 2)

    def test_packaging_requires_current_rights_and_privacy(self):
        source = self.root / 'code.py'; source.write_text('print(1)')
        manifest = {'version': 1, 'files': [{'source': 'code.py', 'destination': 'scripts/code.py'}]}
        with patch.object(packaging, 'ROOT', self.root):
            output = self.root / 'release/hwpdoc'
            result = packaging.build(manifest, {}, output)
            self.assertEqual(result['status'], 'blocked'); self.assertFalse(output.exists())
            # Synthetic grants are local fixtures only, never actual release evidence.
            review = {k: 'SYNTHETIC' for k in ('license', 'source', 'permission_record', 'reviewed_by', 'reviewed_at', 'privacy_review')}
            review.update(sha256=h.sha(source), privacy_sha256=h.sha(source))
            rights = {'files': {'code.py': review}}
            self.assertEqual(packaging.check_release(manifest, rights)['status'], 'ready_to_package')
            source.write_text('print(2)')
            self.assertEqual(packaging.build(manifest, rights, output)['status'], 'blocked')
            self.assertFalse(output.exists())

    def test_packaging_rejects_operational_sources_and_escape(self):
        with patch.object(packaging, 'ROOT', self.root):
            for source, destination in [('knowledge/original.hwpx', 'original.hwpx'), ('code.py', '../escape.py')]:
                with self.subTest(source=source), self.assertRaises(ValueError):
                    packaging.check_release({'version': 1, 'files': [{'source': source, 'destination': destination}]}, {})

    def test_edit_finalize_determinism_and_format_preservation(self):
        from lxml import etree
        header = b'<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"><hh:charProperties><hh:charPr id="10" height="1200"><hh:bold/></hh:charPr></hh:charProperties></hh:head>'
        section = b'<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"><hp:p><hp:run charPrIDRef="10"><hp:secPr><hp:pagePr><hp:margin top="5670" left="8504"/></hp:pagePr></hp:secPr><hp:t>TEST TITLE</hp:t></hp:run><hp:linesegarray><hp:lineseg textpos="0"/></hp:linesegarray></hp:p><hp:p><hp:run><hp:tbl rowCnt="1" colCnt="1"><hp:pos treatAsChar="1"/><hp:tr><hp:tc><hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="1"/><hp:cellSz width="4000" height="2723"/></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></hs:sec>'
        with ZipFile(self.source, 'w') as archive:
            archive.writestr('mimetype', 'application/hwp+zip')
            archive.writestr('Contents/header.xml', header)
            archive.writestr('Contents/section0.xml', section)
        original = h.sha(self.source)
        outputs = []
        for index in range(2):
            output = self.root / f'edited{index}.hwpx'; outputs.append(output)
            commands = [[str(h.SKILL / 'edit_hwpx.py'), str(self.source), '-o', str(output), '--replace', 'TEST TITLE=NEXT TITLE'],
                        [str(h.SKILL / 'finalize_hwpx.py'), str(output), '--strip-linesegarray']]
            for command in commands:
                result = subprocess.run([sys.executable, '-B', '-X', 'utf8', *command], capture_output=True, text=True, encoding='utf-8', timeout=30)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            with ZipFile(output) as archive:
                self.assertEqual(archive.read('Contents/header.xml'), header)
                after = etree.fromstring(archive.read('Contents/section0.xml')); before = etree.fromstring(section)
                ns = {'hp': 'http://www.hancom.co.kr/hwpml/2011/paragraph'}
                for tag in ('margin', 'tbl'):
                    self.assertEqual(etree.tostring(after.find('.//hp:' + tag, ns)), etree.tostring(before.find('.//hp:' + tag, ns)))
                self.assertEqual(after.find('./hp:p/hp:run', ns).get('charPrIDRef'), '10')
                self.assertEqual(after.find('.//hp:t', ns).text, 'NEXT TITLE')
                self.assertIsNone(after.find('.//hp:linesegarray', ns))
        self.assertEqual(h.sha(outputs[0]), h.sha(outputs[1]))
        self.assertEqual(h.sha(self.source), original)

    def test_related_dependency_follows_configured_directory(self):
        related = self.root / 'school-records'; related.mkdir()
        selected = related / 'record.md'; selected.write_text('SOURCE')
        self.settings['references']['related'] = 'school-records'
        h.write_json(self.root / 'evidence.json', {})
        m = dict(version=2, draft={'path': 'source.hwpx'}, evidence={'record': 'evidence.json'}, sources=[], attachments=[],
                 checks={k: {'applicable': k == 'related'} for k in ('budget_names', 'related', 'timetable')})
        dependencies = h.dependencies(m)
        self.assertEqual(dependencies[str(selected)], h.sha(selected))
        selected.write_text('CHANGED')
        self.assertNotEqual(dependencies[str(selected)], h.sha(selected))

    def test_code_update_invalidates_validation_but_preserves_approval(self):
        manifest = self.root / 'approved.json'
        h.write_json(manifest, {'draft_approval': {'record': 'SYNTHETIC approval fixture, not a user approval'}})
        approved = h.sha(manifest)
        code = str(REPO / 'scripts/hwpdoc.py')
        report = dict(work_file=str(self.source), manifest_file=str(manifest), manifest_sha256=approved,
                      dependencies={code: 'old', str(self.source): h.sha(self.source)}, tool_dependencies={code: 'old'},
                      validated_sha256=h.sha(self.source), stages={stage: {'status': 'pass'} for stage in h.STAGES})
        with patch.object(h, 'tool_dependencies', return_value={code: 'new'}):
            self.assertEqual(h.result_state(report), 'unconfirmed')
            h.refresh_validation_tools(report)
            self.assertNotIn('validated_sha256', report)
            self.assertEqual(report['dependencies'][code], 'new')
            self.assertTrue(report['tool_updates'][0]['approval_preserved'])
            self.assertEqual(h.sha(manifest), approved)
            stored = h.report_record(report)
            self.assertEqual(h.report_record(stored, decode=True)['tool_dependencies'], {code: 'new'})

    def test_code_update_cannot_reapprove_changed_source(self):
        manifest = self.root / 'approved.json'; h.write_json(manifest, {})
        code = str(REPO / 'scripts/hwpdoc.py')
        report = dict(work_file=str(self.root / 'work.hwpx'), manifest_file=str(manifest), manifest_sha256=h.sha(manifest),
                      dependencies={code: 'old', str(self.source): h.sha(self.source)}, tool_dependencies={code: 'old'})
        self.source.write_bytes(b'changed input')
        before = copy.deepcopy(report)
        with patch.object(h, 'tool_dependencies', return_value={code: 'new'}), self.assertRaises(ValueError):
            h.refresh_validation_tools(report)
        self.assertEqual(report, before)

    def test_hook_update_and_removal_preserve_user_settings_and_files(self):
        target = self.root / '.codex/hooks.json'
        user_hook = {'matcher': 'Write', 'hooks': [{'type': 'command', 'command': 'user-owned-check'}]}
        h.write_json(target, {'hooks': {'PreToolUse': [user_hook]}, 'user_option': True})
        args = argparse.Namespace(app='codex', tool=['Write'], apply=True, remove=False)
        setup.setup_hooks(args)
        old = h.read_json(target)
        with patch.object(h, 'CODE_ROOT', self.root / 'updated-install'):
            result = setup.setup_hooks(args)
        updated = h.read_json(target)
        self.assertIn(user_hook, updated['hooks']['PreToolUse'])
        self.assertTrue(updated['user_option'])
        self.assertNotEqual(old, updated)
        self.assertEqual(len(updated['hooks']['PreToolUse']), 2)
        self.assertEqual(h.read_json(result['backup'])['settings'], old)
        original = h.sha(self.source)
        args.remove = True
        result = setup.setup_hooks(args)
        self.assertEqual(result['status'], 'removed')
        self.assertEqual(h.read_json(target), {'hooks': {'PreToolUse': [user_hook]}, 'user_option': True})
        self.assertEqual(h.sha(self.source), original)
        self.assertEqual(h.read_json(self.root / '.hwpdoc/hooks-codex-owned.json')['hooks'], {})

    def test_changed_managed_hook_is_not_overwritten(self):
        args = argparse.Namespace(app='claude', tool=None, apply=True, remove=False)
        setup.setup_hooks(args)
        target = self.root / '.claude/settings.json'; changed = h.read_json(target)
        changed['hooks']['PreToolUse'][0]['hooks'][0]['timeout'] = 500
        h.write_json(target, changed)
        with self.assertRaises(ValueError): setup.setup_hooks(args)
        self.assertEqual(h.read_json(target), changed)

    def test_hook_remove_proposal_does_not_apply(self):
        args = argparse.Namespace(app='codex', tool=None, apply=True, remove=False)
        setup.setup_hooks(args)
        target = self.root / '.codex/hooks.json'; original = h.sha(target)
        args.apply = False; args.remove = True
        result = setup.setup_hooks(args)
        self.assertEqual(result['status'], 'proposal')
        self.assertEqual(h.read_json(result['path'])['hooks'], {})
        self.assertEqual(h.sha(target), original)

    def test_registration_visual_artifacts_checked_again_after_registration(self):
        pdf = self.root / 'review.pdf'; pdf.write_bytes(b'SYNTHETIC PDF')
        png = self.root / 'review.png'; png.write_bytes(b'SYNTHETIC PNG')
        r = dict(provenance='converted', purposes=['notice'], hashes={'template.hwpx': h.sha(self.source)})
        review = dict(confirmed=True, by='SYNTHETIC_FIXTURE', at='2026-09-13', record='Not real approval',
                      provenance=r['provenance'], purposes=r['purposes'], hashes=r['hashes'],
                      visual={'record': 'Not actual visual QA', 'document_sha256': h.sha(self.source),
                              'pdf': {'path': 'review.pdf', 'sha256': h.sha(pdf)},
                              'images': [{'path': 'review.png', 'sha256': h.sha(png)}]})
        self.assertEqual(t.review_errors(r, review), [])
        png.write_bytes(b'MODIFIED')
        self.assertTrue(any('해시 변경' in reason for reason in t.review_errors(r, review)))
        review['visual']['images'][0]['path'] = '../outside.png'
        self.assertTrue(t.review_errors(r, review))

    def test_legacy_migration_creates_draft_and_detects_source_change(self):
        name = '가정통신문'
        folder = self.root / 'knowledge/templates' / name; folder.mkdir(parents=True)
        shutil.copy2(self.source, folder / '양식.hwpx')
        (folder / 'README.md').write_text('Synthetic legacy fixture', encoding='utf-8')
        for filename in ('slots.json', 'budget.json', 'structure.json'): h.write_json(folder / filename, {})
        with patch.object(h, 'TEMPLATES', (name,)):
            result = t.migrate_legacy()[0]
            registration = self.root / result['registration']
            self.assertEqual(result['registration_status'], 'draft')
            self.assertEqual(h.sha(registration / 'original.hwpx'), h.sha(self.source))
            self.assertEqual(h.read_json(registration / 'mapping.proposed.json'), {'p:0': 'p:0'})
            self.assertFalse((registration / 'review.json').exists())
            original_record = h.sha(registration / 'registration.json')
            t.migrate_legacy()
            self.assertEqual(h.sha(registration / 'registration.json'), original_record)
            (folder / 'README.md').write_text('Changed source')
            self.assertEqual(t.migrate_legacy()[0]['status'], 'source_changed')

    def test_review_draft_does_not_invent_human_confirmation(self):
        folder = t.folder_for('review-draft', '1'); folder.mkdir(parents=True)
        for name in t.FILES: (folder / name).write_bytes(b'SYNTHETIC')
        record = {'status': 'awaiting_review', 'hashes': t.file_hashes(folder), 'purposes': ['test'], 'provenance': 'native'}
        result = t.prepare_review(folder, record)
        draft = h.read_json(result['review_draft'])
        self.assertIs(draft['confirmed'], False)
        self.assertEqual([draft[key] for key in ('by', 'at', 'record')], ['', '', ''])
        self.assertEqual(draft['hashes'], record['hashes'])
        self.assertEqual(result['status'], 'awaiting_review')
        draft['record'] = 'Human may edit this later'
        h.write_json(result['review_draft'], draft)
        t.prepare_review(folder, record)
        self.assertEqual(h.read_json(result['review_draft'])['record'], draft['record'])

    def test_pdf_preview_covers_all_pages_without_visual_approval(self):
        import fitz
        import hwpdoc_preview
        pdf = self.root / 'synthetic.pdf'
        with fitz.open() as document:
            for _ in range(2): document.new_page().insert_text((40, 40), 'Synthetic developer fixture')
            document.save(pdf)
        result = hwpdoc_preview.render(pdf, self.root / 'preview', self.ctx)
        self.assertEqual(len(result['images']), 2)
        self.assertEqual(result['visual_review'], 'unconfirmed')
        self.assertEqual(result, hwpdoc_preview.render(pdf, self.root / 'preview', self.ctx))
        first_image = self.root / result['images'][0]['path']; first_image.write_bytes(b'changed image')
        with self.assertRaises(ValueError): hwpdoc_preview.render(pdf, self.root / 'preview', self.ctx)

    def test_report_roundtrip_preserves_approval_and_move(self):
        report = dict(work_file=str(self.source), manifest_file=str(self.root / 'manifest.json'),
                      dependencies={str(self.source): h.sha(self.source)}, stages={'edit': {'command': [str(REPO / 'scripts/hwpdoc.py'), str(self.source), '--flag']}}, history=[],
                      delivery={'local': str(self.root / 'output/result.hwpx')}, copy={})
        original = copy.deepcopy(report)
        stored = h.report_record(report)
        self.assertEqual(original, report)
        self.assertEqual(stored['work_file'], 'workspace:source.hwpx')
        self.assertEqual(h.report_record(stored, decode=True)['work_file'], str(self.source))
        self.assertEqual(stored['stages']['edit']['command'][0], 'plugin:scripts/hwpdoc.py')

    def test_unknown_stage_cannot_pass(self):
        m = self.root / 'manifest.json'; h.write_json(m, {})
        report = dict(work_file=str(self.source), manifest_file=str(m), manifest_sha256=h.sha(m), dependencies={},
                      validated_sha256=h.sha(self.source), runtime=h.runtime_snapshot(), stages={s: {'status': 'pass'} for s in h.STAGES})
        self.assertEqual(h.result_state(report), 'pass')
        report['stages']['hancom']['status'] = 'invented'
        self.assertEqual(h.result_state(report), 'unconfirmed')

    def test_registration_draft_is_not_available(self):
        args = self.registration_args(); args.source = str(self.source)
        before = h.sha(self.source)
        r = t.add_template(args)
        self.assertEqual(r['status'], 'draft')
        self.assertEqual(h.sha(self.source), before)
        with self.assertRaises(ValueError): t.resolve_template({'id': args.id, 'version': args.version, 'sha256': before})
        with self.assertRaises(ValueError): t.add_template(args)

    def test_invalid_mapping_and_twice_failure_stop(self):
        args = self.registration_args(); args.source = str(self.source); t.add_template(args)
        args.source = None; args.mapping = str(self.root / 'mapping.json')
        h.write_json(args.mapping, {'title': 'p:999'})
        with self.assertRaises(ValueError): t.add_template(args)
        h.write_json(args.mapping, {'title': 'p:0'})
        with patch.object(h, 'run', return_value=(1, 'synthetic schema failure')) as run:
            for _ in range(2):
                result = t.add_template(args)
                self.assertEqual(result['stages']['structure']['status'], 'fail')
            with self.assertRaises(ValueError): t.add_template(args)
            self.assertEqual(run.call_count, 2)

    def test_template_cli_internal_failure_nonzero(self):
        r = self.cli('add-template', '--id', 'invalid', '--version', '1', '--source', str(self.source), '--purpose', 'notice')
        self.assertEqual(r.returncode, 0, r.stderr)
        h.write_json(self.root / 'mapping.json', {'title': 'p:0'})
        r = self.cli('add-template', '--id', 'invalid', '--version', '1', '--mapping', 'mapping.json')
        self.assertNotEqual(r.returncode, 0)
        self.assertNotEqual(json.loads(r.stdout)['status'], 'registered')

    def test_prepare_leaves_judgments_unknown_and_imports_source(self):
        args = argparse.Namespace(job='prepare', template=None, version=None, source=[str(self.source)], values=None, workflow='W2', document_type='reply')
        result = w.prepare(args)
        m = h.read_json(result['manifest'])
        self.assertFalse(m['draft_approval']['confirmed'])
        self.assertIsNone(m['received']['submission_method'])
        self.assertTrue(all(x['applicable'] is None for x in m['checks'].values()))
        self.assertFalse(Path(m['sources'][0]['path']).is_absolute())
        self.assertEqual(h.sha(self.root / m['sources'][0]['path']), h.sha(self.source))
        self.assertEqual(m['evidence']['record'], '')

    def test_prepare_legacy_missing_values_and_explicit_blank(self):
        for workflow in ('W1', 'W2'):
            args = argparse.Namespace(job='legacy-' + workflow, template='legacy', version=None,
                                      source=[str(self.source)], values=None, workflow=workflow, document_type='notice')
            with patch.object(h, 'template_file', return_value=self.source):
                result = w.prepare(args)
            m = h.read_json(result['manifest'])
            self.assertEqual(m['slots']['values'], {})
            self.assertTrue(any('값 미입력: p:0' in x for x in m['unresolved']))
            self.assertEqual(result['missing'], m['unresolved'])
            self.assertIn('값 미입력: p:0', Path(result['manifest']).with_name('확인할사항.md').read_text(encoding='utf-8'))
            self.assertFalse(m['draft_approval']['confirmed'])
            if workflow == 'W2':
                self.assertIsNone(m['received']['submission_method'])
        h.write_json(self.root / 'blank.json', {'p:0': ''})
        args.job = 'intentional-blank'; args.values = str(self.root / 'blank.json')
        with patch.object(h, 'template_file', return_value=self.source):
            result = w.prepare(args)
        self.assertFalse(any('값 미입력:' in x for x in result['missing']))

    def test_absent_school_data_only_blocks_its_command(self):
        r = self.cli('list-templates'); self.assertEqual(r.returncode, 0, r.stderr)
        env = dict(os.environ, HWPDOC_WORKSPACE=str(self.root), PYTHONDONTWRITEBYTECODE='1')
        r = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(REPO / 'scripts/budget_name_check.py'), '--list'], env=env, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(r.returncode, 2)
        self.assertIn('예산 자료 없음', r.stderr)

    def test_other_project_protection_scope_and_native_powershell(self):
        sys.path.insert(0, str(REPO / '.claude/hooks'))
        from event_adapter import evaluate
        local = self.root / 'knowledge/file'; local.parent.mkdir(); local.write_text('KEEP')
        other = self.root / 'nested-other'; other.mkdir()
        # A nested unmarked folder inherits this workspace. A marked unrelated
        # project is represented outside this workspace in a separate temp root.
        with tempfile.TemporaryDirectory(dir=self.root.parent) as outside, patch('hwpdoc_config.CODE_ROOT', self.root / 'installation'):
            target = Path(outside) / 'docs/file'; target.parent.mkdir(); target.write_text('KEEP')
            self.assertIsNone(owned_workspace(target))
            result = evaluate({'cwd': str(self.root), 'tool_name': 'Write', 'tool_input': {'file_path': str(target)}})
            self.assertEqual(result['decision'], 'allow')
        result = evaluate({'cwd': str(self.root), 'tool_name': 'PowerShell', 'tool_input': {'command': 'Set-Content knowledge/file changed'}})
        self.assertEqual(result['decision'], 'deny')
        self.assertEqual(local.read_text(), 'KEEP')


if __name__ == '__main__': unittest.main()
